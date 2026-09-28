#!/usr/bin/env python3
"""Call DEPs per dataset and build a gene-by-dataset summary table.

DEP: p < 0.05 and |log2FC| > 1 (both adjustable from the command line).

Usage:
    python make_dep_summary.py GEC_tissue_proteomics_DEP_Master_table_20260917.xlsx
    python make_dep_summary.py input.xlsx -o out.xlsx --p 0.05 --lfc 1
"""
import argparse

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATASETS = ["PG1", "PG2", "PG3", "PG4", "PE1", "PE2", "PE3", "PE5"]

# (gene col, log2FC col, p col) for sheets that only need renaming
COLS = {
    "PG1": ("Gene", "Fold_Change", "P_Value"),
    "PG3": ("Gene", "Fold_Change", "P_Value"),
    "PE5": ("Gene", "Fold_Change", "P_Value"),
    "PG4": ("Gene", "log2FC_tumor_vs_normal", "P_Value"),
    "PE1": ("Gene symbol", "log2FC_tumor_vs_normal", "P_Value"),
}

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
UP_FILL = PatternFill("solid", fgColor="F8D7DA")
DOWN_FILL = PatternFill("solid", fgColor="D6E4F5")
HEADER_FILL = PatternFill("solid", fgColor="D9EAD3")


def uniprot_to_gene(sheets):
    # PG2 has UniProt accessions only, so borrow gene names from other sheets
    pairs = [
        sheets["PG1"][["ID", "Gene"]],
        sheets["PE1"][["UniProt ID", "Gene symbol"]],
        sheets["PE2"][["Protein accession", "Gene name"]],
    ]
    pairs = [p.set_axis(["acc", "gene"], axis=1) for p in pairs]
    m = pd.concat(pairs).dropna().drop_duplicates("acc")
    return m.set_index("acc")["gene"]


def load_dataset(name, sheets, acc2gene):
    df = sheets[name]
    if name in COLS:
        g, fc, p = COLS[name]
        out = pd.DataFrame({"Gene": df[g], "log2FC": df[fc], "P": df[p]})
    elif name == "PG2":
        ids = df["Protein ID"]
        out = pd.DataFrame({
            "Gene": ids.map(acc2gene).fillna(ids),
            "log2FC": df["log2FC_tumor_vs_normal"],
            "P": df["P_Value"],
        })
    elif name == "PE2":
        # Ratio is on a linear scale
        out = pd.DataFrame({
            "Gene": df["Gene name"],
            "log2FC": np.log2(df["Ratio"]),
            "P": df["P value"],
        })
    else:
        raise ValueError(f"no loader for {name}")
    return out.dropna()


def call_deps(df, p_cut, lfc_cut):
    df = df[(df["P"] < p_cut) & (df["log2FC"].abs() > lfc_cut)].copy()
    df["Direction"] = np.where(df["log2FC"] > 0, "Up", "Down")
    # keep the most significant row when a gene appears more than once
    return df.sort_values("P").drop_duplicates("Gene")


def pe3_calls(df):
    # PE3 only reports direction; drop genes whose duplicate rows disagree
    df = df.dropna(subset=["Name", "consistent_direction"])
    agree = df.groupby("Name")["consistent_direction"].nunique() == 1
    df = df[df["Name"].isin(agree[agree].index)].drop_duplicates("Name")
    return pd.DataFrame({
        "Gene": df["Name"],
        "log2FC": np.nan,
        "P": np.nan,
        "Direction": df["consistent_direction"],
    })


def summarize(dep_tables):
    calls = {k: v.set_index("Gene")["Direction"] for k, v in dep_tables.items()}
    summ = pd.DataFrame(calls)[DATASETS]
    n_up = (summ == "Up").sum(axis=1)
    n_down = (summ == "Down").sum(axis=1)

    consistent = pd.Series("", index=summ.index)
    consistent[(n_up > 0) & (n_down == 0)] = "Up"
    consistent[(n_down > 0) & (n_up == 0)] = "Down"

    order = pd.DataFrame({
        "n": n_up + n_down,
        "top": np.maximum(n_up, n_down),
        "dir": consistent.map({"Down": 0, "Up": 1, "": 2}),
        "gene": summ.index,
    })
    order = order.sort_values(["n", "top", "dir", "gene"],
                              ascending=[False, False, True, True])
    return summ.loc[order.index].rename_axis("Gene").reset_index()


def style_header(ws):
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")


def write_summary(wb, summ):
    ws = wb.active
    ws.title = "DEP_summary"
    ws.append(["Gene"] + DATASETS + ["n_Up", "n_Down", "n_datasets", "consistent_"])

    first = get_column_letter(2)
    last = get_column_letter(len(DATASETS) + 1)
    up_col, down_col = get_column_letter(len(DATASETS) + 2), get_column_letter(len(DATASETS) + 3)

    for r, row in enumerate(summ.itertuples(index=False), start=2):
        cells = [v if isinstance(v, str) else None for v in row]
        up, down = f"{up_col}{r}", f"{down_col}{r}"
        cells += [
            f'=COUNTIF({first}{r}:{last}{r},"Up")',
            f'=COUNTIF({first}{r}:{last}{r},"Down")',
            f"={up}+{down}",
            f'=IF(AND({up}>0,{down}=0),"Up",IF(AND({down}>0,{up}=0),"Down",""))',
        ]
        ws.append(cells)

    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            if c.value == "Up":
                c.fill = UP_FILL
            elif c.value == "Down":
                c.fill = DOWN_FILL
    style_header(ws)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions[get_column_letter(ws.max_column)].width = 13


def write_dep_list(wb, dep_tables):
    ws = wb.create_sheet("DEP_list_by_dataset")
    df = pd.concat([t.assign(Dataset=k) for k, t in dep_tables.items()])
    df = df[["Dataset", "Gene", "log2FC", "P", "Direction"]]
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
    style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_criteria(wb, p_cut, lfc_cut):
    ws = wb.create_sheet("Criteria")
    rows = [
        ("Item", "Value"),
        ("P value cutoff", f"P < {p_cut} (raw p-value)"),
        ("Fold change cutoff", f"|log2FC| > {lfc_cut}"),
        ("PE2", "Ratio converted with log2(Ratio); rows without a p-value removed"),
        ("PE3", "No FC/p-value in source; consistent_direction used directly. "
                "Genes with conflicting duplicate rows removed"),
        ("PG2", "UniProt accessions mapped to genes using PG1/PE1/PE2; "
                "unmapped accessions kept as-is"),
        ("Duplicate genes", "Row with the smallest p-value kept"),
        ("consistent_", "Up/Down when all calls agree, blank when mixed"),
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
    style_header(ws)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="master table (.xlsx), one sheet per dataset")
    ap.add_argument("-o", "--output", default="GEC_DEP_summary.xlsx")
    ap.add_argument("--p", type=float, default=0.05, help="p-value cutoff (default 0.05)")
    ap.add_argument("--lfc", type=float, default=1.0, help="|log2FC| cutoff (default 1)")
    args = ap.parse_args()

    sheets = pd.read_excel(args.input, sheet_name=None)
    acc2gene = uniprot_to_gene(sheets)

    deps = {}
    for name in DATASETS:
        if name == "PE3":
            deps[name] = pe3_calls(sheets[name])
        else:
            df = load_dataset(name, sheets, acc2gene)
            deps[name] = call_deps(df, args.p, args.lfc)

    summ = summarize(deps)

    wb = Workbook()
    write_summary(wb, summ)
    write_dep_list(wb, deps)
    write_criteria(wb, args.p, args.lfc)
    wb.save(args.output)

    print("DEPs per dataset:")
    for k, v in deps.items():
        print(f"  {k}: {len(v)}")
    print(f"{len(summ)} genes written to {args.output}")
    n_unmapped = (~sheets["PG2"]["Protein ID"].isin(acc2gene.index)).sum()
    if n_unmapped:
        print(f"note: {n_unmapped} PG2 accessions had no gene name")


if __name__ == "__main__":
    main()
